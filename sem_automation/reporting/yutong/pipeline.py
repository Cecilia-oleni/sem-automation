# 模块：sem_automation/reporting/yutong/pipeline.py；内部模块由统一入口调用。
# VS Code PowerShell 先输入：Set-Location -LiteralPath 'D:\sem自动化 - 副本'
# 终端输入（复制时去掉注释符）：& '.\.venv\Scripts\python.exe' -X utf8 '.\sem.py' reports yutong --config '.\config\reports\yutong\2026-08.json' --month 2026-08
# 该示例复用本地数据；详见 docs/月报操作说明.md。
"""Read-only input validation, staged Yandex report rendering, candidate manifest."""
import json
import math
import re
import calendar
from pathlib import Path
import subprocess
import shutil
from datetime import date, datetime
from openpyxl import load_workbook
from sem_automation.core.paths import PROJECT_ROOT, renderer
from sem_automation.core.runs import new_run, digest, write_json
from sem_automation.reporting.standard.pipeline import resolve_node
from sem_automation.reporting.yutong.history import periods, blocks, require_periods

COUNTRIES=['俄罗斯','阿塞拜疆','亚美尼亚','摩尔多瓦','格鲁吉亚','白俄罗斯','哈萨克斯坦','乌兹别克斯坦','吉尔吉斯斯坦','塔吉克斯坦','土库曼斯坦']
PATHS=['searchNetworkCurrent','searchNetworkPriorMonth','searchNetworkPriorYear','leadSource','countrySource','monthlyWorkbook','countryHistory','ytdHistory']

def workbook_rows(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        for sheet in wb:
            sheet.reset_dimensions()
        return {s.title:[[v.strftime('%Y-%m') if isinstance(v,(date,datetime)) else v for v in row] for row in s.iter_rows(values_only=True)] for s in wb}
    finally:wb.close()

def workbook_formulas(path):
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        result={}
        for sheet in wb:
            sheet.reset_dimensions()
            result[sheet.title]={cell.coordinate:cell.value for row in sheet.iter_rows() for cell in row if cell.data_type=='f'}
        return result
    finally:wb.close()

def error_cells(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        found={}
        for sheet in wb:
            sheet.reset_dimensions()
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type=='e':found[f'{sheet.title}!{cell.coordinate}']=cell.value
        return found
    finally:wb.close()

def verify_outputs(config,stage):
    month=config['reportMonth'];known={}
    for key,filename in [('monthlyWorkbook',f'yutong_{month}_monthly_country_leads.xlsx'),('countryHistory',f'yutong_country_history_through_{month}.xlsx'),('ytdHistory',f'yutong_ytd_history_through_{month}.xlsx')]:
        original=error_cells(config['paths'][key]);actual=error_cells(stage/filename)
        introduced={cell:error for cell,error in actual.items() if original.get(cell)!=error}
        if introduced:raise ValueError(f'{filename} 生成了新的 Excel 错误: {introduced}')
        known[key]=actual
    ytd=workbook_rows(stage/f'yutong_ytd_history_through_{month}.xlsx')
    index=blocks(ytd['月度源数据'],COUNTRIES,'候选YTD月度源数据')
    require_periods(index,config['_periods']['ytd'],'候选YTD')
    for country,lead_row,spend_row in zip(COUNTRIES,[5,6,7,8,9,10,12,13,14,15,16],[23,24,25,26,27,28,30,31,32,33,34]):
        for column,target in [(5,lead_row),(4,spend_row)]:
            values=[index[m]['rows'][country][column] for m in config['_periods']['ytd']]
            if any(not isinstance(v,(int,float)) and v!='/' for v in values):raise ValueError(f'YTD 缺少数值: {country}')
            expected=sum(v for v in values if isinstance(v,(int,float)))
            for sheet in ['国家YTD汇总_数值','国家YTD汇总_公式']:
                actual=ytd[sheet][target-1][1]
                if not isinstance(actual,(int,float)) or not math.isclose(actual,expected,rel_tol=1e-9,abs_tol=1e-6):raise ValueError(f'{sheet}: {country} YTD 与逐月汇总不一致')
    return {'new_excel_errors':0,'preserved_source_errors':known,'ytd_recomputed':True}

def prepare(config_path,month):
    config_path=Path(config_path).resolve()
    config=json.loads(config_path.read_text(encoding='utf-8'))
    span=periods(month)
    if config.get('client')!='yutong' or config.get('reportMonth')!=month:raise ValueError('客户必须为 yutong，月份必须与周期配置一致')
    rate=config.get('exchangeRateCnyPerUsd')
    if not isinstance(rate,(int,float)) or not math.isfinite(rate) or rate<=0:raise ValueError('汇率必须为正数')
    for name in ('current','priorMonth','priorYear'):
        value=config.get('leadComparisonTotals',{}).get(name)
        if not isinstance(value,int) or isinstance(value,bool) or value<0:raise ValueError(f'人工核对线索数 {name} 必须为非负整数')
    if not config.get('sources',{}).get('exchangeRate') or not config.get('sources',{}).get('leadComparisonTotals'):
        raise ValueError('配置 sources 必须说明汇率和人工核对总数的来源')
    paths={k:(config_path.parent/Path(config['paths'][k])).resolve() for k in PATHS}
    missing=[f'{k}: {p}' for k,p in paths.items() if not p.is_file()]
    if missing:raise FileNotFoundError('缺少输入文件：\n'+'\n'.join(missing))
    if len(set(paths.values()))!=len(paths):raise ValueError('不同期间/用途不能引用同一个输入文件')
    hashes={k:digest(p) for k,p in paths.items()}
    books={k:workbook_rows(p) for k,p in paths.items()}
    for key,period in [('searchNetworkCurrent',month),('countrySource',month),('searchNetworkPriorMonth',span['previous']),('searchNetworkPriorYear',span['yearAgo'])]:
        rows=books[key].get('Sheet0',[])
        text=' '.join(str(v) for row in rows[:5] for v in row if v is not None)
        year,number=map(int,period.split('-'))
        dates=re.findall(r'\d{2}\.\d{2}\.\d{4}',text)
        expected=[f'01.{number:02d}.{year}',f'{calendar.monthrange(year,number)[1]:02d}.{number:02d}.{year}']
        if dates!=expected:raise ValueError(f'{key}: 导出期间 {dates} 与 {period} 不一致')
        if config.get('directLogin') and config['directLogin'] not in text:raise ValueError(f'{key}: 客户账号不一致')
    for key,sheets in [('countryHistory',['整理数值','原始数据','公式溯源']),('ytdHistory',['月度源数据','国家YTD汇总_数值','国家YTD汇总_公式']),('monthlyWorkbook',['Sheet1']),('leadSource',['Sheet1']),('countrySource',['Sheet0'])]:
        for name in sheets:
            if name not in books[key]:raise ValueError(f'{key}: 缺少工作表 {name}')
    histories={}
    canonical=blocks(books['countryHistory']['整理数值'],COUNTRIES,'countryHistory/整理数值')
    for row in books['countryHistory']['原始数据']:
        if row and isinstance(row[0],(int,float)) and str(row[0]).endswith('.1'):
            year=str(row[0]).split('.')[0]
            choices=[k for k in (year+'-01',year+'-10') if k in canonical]
            if len(choices)!=1:raise ValueError(f'原始数据月份 {row[0]} 有歧义，请改为 YYYY-MM')
            row[0]=choices[0]
    for key,names in [('countryHistory',['整理数值','原始数据','公式溯源']),('ytdHistory',['月度源数据'])]:
        histories[key]={}
        for name in names:
            index=blocks(books[key][name],COUNTRIES,f'{key}/{name}',raw=name=='原始数据')
            require_periods(index,[span['previous'],span['yearAgo']]+span['ytd'][:-1],f'{key}/{name}')
            histories[key][name]=index
    for key,p in paths.items():
        if digest(p)!=hashes[key]:raise ValueError(f'读取期间源文件发生变化: {p}')
    config['paths']={k:str(p) for k,p in paths.items()}
    config['_periods']=span;config['_books']=books;config['_histories']=histories
    config['_formulas']={key:workbook_formulas(paths[key]) for key in ['monthlyWorkbook','countryHistory','ytdHistory']}
    return config,hashes

def run_report(*,config_path,month,node=None,check_only=False):
    config,hashes=prepare(config_path,month)
    if check_only:return '输入结构和历史期间校验通过（未写文件）'
    runtime=resolve_node(node)
    run=new_run('yutong','yandex',month)
    stage=run/'_internal/staging';stage.mkdir()
    input_config=run/'_internal/resolved_config.json'
    write_json(input_config,config)
    try:
        subprocess.run([runtime,str(renderer('yutong','country_leads.mjs')),str(input_config),str(stage)],cwd=PROJECT_ROOT,check=True)
        monthly=stage/f'yutong_{month}_monthly_country_leads.xlsx'
        subprocess.run([runtime,str(renderer('yutong','search_network.mjs')),str(input_config),str(monthly)],cwd=PROJECT_ROOT,check=True)
        write_json(run/'_internal/workbook_checks.json',verify_outputs(config,stage))
        entries=[]
        for key,filename in [('countryHistory',f'yutong_country_history_through_{month}.xlsx'),('ytdHistory',f'yutong_ytd_history_through_{month}.xlsx')]:
            candidate=run/'history_candidates'/filename
            shutil.move(str(stage/filename),candidate)
            entries.append({'source':config['paths'][key],'source_sha256':hashes[key],'candidate':str(candidate.resolve()),'candidate_sha256':digest(candidate)})
        for key,path in config['paths'].items():
            if digest(path)!=hashes[key]:raise ValueError(f'生成期间输入发生变化: {path}')
        shutil.move(str(monthly),run/'deliverables'/monthly.name)
        for p in stage.glob('*.json'):shutil.move(str(p),run/'_internal'/p.name)
        manifest={'schema':'yutong_sync_v1','status':'completed','client':'yutong','month':month,'run':str(run.resolve()),'entries':entries,'inputs':hashes}
        write_json(run/'_internal/manifest.json',manifest)
        write_json(run/'_internal/status.json',{'status':'completed'})
        return run
    except Exception as exc:
        write_json(run/'_internal/status.json',{'status':'failed','error':str(exc)})
        raise
