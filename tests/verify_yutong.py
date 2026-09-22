# 终端输入：& '.\.venv\Scripts\python.exe' -X utf8 '.\tests\verify_yutong.py'
# 只读验收：新八月工作簿与历史已验证交付对比；确认 YTD 从月份源数据重算。
from pathlib import Path
import sys
import math
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from sem_automation.reporting.yutong.pipeline import COUNTRIES, workbook_rows
from sem_automation.reporting.yutong.history import blocks
root=Path(__file__).resolve().parents[1]
run=sorted((root/'outputs/reports/yutong/yandex/2026-08/runs').glob('*/_internal/manifest.json'))[-1].parent.parent
old_path=root/'outputs/_archive/yutong_monthly_2026-08_run_001/yutong_2026-08_monthly_country_leads.xlsx'
old=load_workbook(old_path,data_only=True)['Sheet1'];new=load_workbook(next((run/'deliverables').glob('*.xlsx')),data_only=True)['Sheet1']
diff=[]
for area in ['A1:O16','A21:M34','A37:R54','A151:M169']:
    for row in old[area]:
        for c in row:
            v=new[c.coordinate].value
            if isinstance(c.value,(int,float)) and isinstance(v,(int,float)) and not math.isclose(c.value,v,abs_tol=1e-6,rel_tol=1e-9):diff.append((c.coordinate,c.value,v))
print('Numeric changes from baseline:',diff)
ytd=next((run/'history_candidates').glob('*ytd_history*.xlsx'))
books=workbook_rows(ytd);history=blocks(books['月度源数据'],COUNTRIES)
assert len([k for k in history if k=='2026-08'])==1
for i,name in enumerate(COUNTRIES):
    values=[history[f'2026-{month:02d}']['rows'][name][5] for month in range(1,9)]
    expected=sum(v for v in values if isinstance(v,(int,float)))
    target=[5,6,7,8,9,10,12,13,14,15,16][i]
    assert books['国家YTD汇总_数值'][target-1][1]==expected,(name,expected,books['国家YTD汇总_数值'][target-1][1])
    assert books['国家YTD汇总_公式'][target-1][1]==expected
errors=[]
for p in list((run/'deliverables').glob('*.xlsx'))+list((run/'history_candidates').glob('*.xlsx')):
    wb=load_workbook(p,data_only=True)
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.data_type=='e':errors.append((p.name,sheet.title,cell.coordinate,cell.value))
print('Excel errors:',errors[:12], 'total',len(errors))
print('YTD source recalculation and same-month uniqueness: OK')
