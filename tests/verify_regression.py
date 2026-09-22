# 终端输入：& '.\.venv\Scripts\python.exe' -X utf8 '.\tests\verify_regression.py'
# 只读比较已有验收文件与新生成文件，不写工作簿。
from pathlib import Path
import json
import math
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]

def latest(client,kind,period):
    return sorted((ROOT/f'outputs/reports/{client}/{kind}/{period}/runs').glob('*/deliverables'),key=lambda p:p.parent.name)[-1]

def compare(old,new):
    a=load_workbook(old,data_only=False);b=load_workbook(new,data_only=False)
    differences=[]
    if a.sheetnames!=b.sheetnames:differences.append(('sheets',a.sheetnames,b.sheetnames))
    for name in a.sheetnames:
        if name not in b:continue
        x,y=a[name],b[name]
        if str(x.merged_cells)!=str(y.merged_cells):differences.append((name,'merged'))
        if len(x._charts)!=len(y._charts):differences.append((name,'charts'))
        if str(x.print_area)!=str(y.print_area):differences.append((name,'print_area'))
        for row in x:
            for c in row:
                before,after=c.value,y[c.coordinate].value
                if isinstance(before,(int,float)) and isinstance(after,(int,float)):
                    same=math.isclose(before,after,rel_tol=1e-10,abs_tol=1e-7)
                else:same=before==after
                if not same:differences.append((name,c.coordinate,before,after))
    a.close();b.close()
    return differences

if __name__=='__main__':
    result={}
    for p in latest('angel_yeast','angel-yeast','2026-08').glob('*.xlsx'):
        result[p.name]=compare(ROOT/'outputs/_archive/angel_yeast/2026-08'/p.name,p)
    for p in latest('lingyu','standard','2026-07').glob('*.xlsx'):
        result[p.name]=compare(ROOT/'outputs/_archive/lingyu/2026-07/report_classic'/p.name,p)
    target=ROOT/'outputs/_migration/2026-09-20/regression.json'
    target.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print({k:len(v) for k,v in result.items()})
    print('Differences:',target)
